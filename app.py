from flask import Flask, render_template, request, jsonify, send_file
from datetime import datetime
import json, os, io
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)

MASTER_DATA = {
    "10.1": {"name": "ประชาอุทิศ", "area": "สมุทร (007)", "type": "ป้ายโฆษณา"},
    "10.2": {"name": "สุขุมวิท 101", "area": "กรุงเทพ (001)", "type": "ป้ายโฆษณา"},
    "10.3": {"name": "ลาดพร้าว 87", "area": "กรุงเทพ (001)", "type": "ป้ายโฆษณา"},
    "11.1": {"name": "รามคำแหง 24", "area": "กรุงเทพ (002)", "type": "ป้ายโฆษณา"},
    "11.2": {"name": "บางนา กม.3", "area": "กรุงเทพ (003)", "type": "ป้ายโฆษณา"},
    "12.1": {"name": "อนุสาวรีย์ชัย", "area": "กรุงเทพ (004)", "type": "ป้ายโฆษณา"},
    "12.2": {"name": "ดอนเมือง ขาออก", "area": "กรุงเทพ (005)", "type": "ป้ายโฆษณา"},
    "13.1": {"name": "เซ็นทรัล ลาดพร้าว", "area": "กรุงเทพ (006)", "type": "ป้ายโฆษณา"},
}

TRAINING_MATRIX = {
    "ไม่ใส่ PPE": {"level": "สูง", "course": "PPE Safety", "type": "New Skill", "duration": "2 ชั่วโมง"},
    "Ground เกินมาตรฐาน": {"level": "สูง", "course": "Grounding System", "type": "Reskill", "duration": "3 ชั่วโมง"},
    "Continuity เกินมาตรฐาน": {"level": "ปานกลาง", "course": "Electrical Continuity", "type": "Upskill", "duration": "2 ชั่วโมง"},
    "Voltage ไม่ถูกต้อง": {"level": "สูง", "course": "Voltage System", "type": "Upskill", "duration": "3 ชั่วโมง"},
    "Location ไม่ตรง": {"level": "วิกฤต", "course": "Location Audit", "type": "Upskill", "duration": "1 ชั่วโมง"},
    "อุปกรณ์ไม่ผ่าน": {"level": "ปานกลาง", "course": "Equipment Inspection", "type": "Reskill", "duration": "2 ชั่วโมง"},
}

inspection_records = []


def validate_inspection(data):
    issues = []
    result = "PASS"
    training_triggers = []

    # Step 2: PPE check
    ppe_items = data.get("ppe", [])
    required_ppe = ["helmet", "gloves", "boots", "tools"]
    missing_ppe = [p for p in required_ppe if p not in ppe_items]
    if missing_ppe:
        issues.append(f"PPE ไม่ครบ: {', '.join(missing_ppe)}")
        result = "STOP"
        training_triggers.append("ไม่ใส่ PPE")

    # Step 3: Electrical values
    try:
        ground = float(data.get("ground", 999))
        if ground > 5:
            issues.append(f"Ground = {ground}Ω (เกิน 5Ω)")
            result = "FAIL" if result != "STOP" else result
            training_triggers.append("Ground เกินมาตรฐาน")
    except (ValueError, TypeError):
        issues.append("ค่า Ground ไม่ถูกต้อง")

    try:
        continuity = float(data.get("continuity", 999))
        if continuity >= 0.5:
            issues.append(f"Continuity = {continuity}Ω (เกิน 0.5Ω)")
            result = "FAIL" if result not in ["STOP"] else result
            training_triggers.append("Continuity เกินมาตรฐาน")
    except (ValueError, TypeError):
        issues.append("ค่า Continuity ไม่ถูกต้อง")

    try:
        voltage = float(data.get("voltage", 0))
        if voltage not in [220, 380] and not (215 <= voltage <= 225) and not (375 <= voltage <= 385):
            issues.append(f"Voltage = {voltage}V (ต้องเป็น 220 หรือ 380V)")
            result = "FAIL" if result not in ["STOP"] else result
            training_triggers.append("Voltage ไม่ถูกต้อง")
    except (ValueError, TypeError):
        issues.append("ค่า Voltage ไม่ถูกต้อง")

    # Step 4: Equipment
    equipment_items = data.get("equipment", [])
    required_equipment = ["breaker", "timer", "magnetic", "wiring"]
    missing_equipment = [e for e in required_equipment if e not in equipment_items]
    if missing_equipment:
        issues.append(f"อุปกรณ์ไม่ผ่าน: {', '.join(missing_equipment)}")
        if result not in ["STOP", "CRITICAL"]:
            result = "FAIL"
        training_triggers.append("อุปกรณ์ไม่ผ่าน")

    # Step 5: Location check
    code_id = data.get("code_id", "").strip()
    if code_id not in MASTER_DATA:
        issues.append(f"Location Code '{code_id}' ไม่ตรง Master Data")
        result = "CRITICAL"
        training_triggers.append("Location ไม่ตรง")

    # Training recommendations
    training_recs = []
    seen = set()
    for trigger in training_triggers:
        if trigger in TRAINING_MATRIX and trigger not in seen:
            seen.add(trigger)
            training_recs.append({
                "issue": trigger,
                **TRAINING_MATRIX[trigger]
            })

    return {
        "result": result,
        "issues": issues,
        "training": training_recs,
        "location_info": MASTER_DATA.get(code_id, {})
    }


@app.route("/")
def index():
    return render_template("index.html", master_data=MASTER_DATA)


@app.route("/api/inspect", methods=["POST"])
def inspect():
    data = request.json
    validation = validate_inspection(data)

    record = {
        "id": len(inspection_records) + 1,
        "date": datetime.now().strftime("%d/%m/%Y %H:%M"),
        "technician": data.get("technician", ""),
        "code_id": data.get("code_id", ""),
        "location": data.get("location", ""),
        "ground": data.get("ground", ""),
        "continuity": data.get("continuity", ""),
        "voltage": data.get("voltage", ""),
        "ppe": data.get("ppe", []),
        "equipment": data.get("equipment", []),
        "result": validation["result"],
        "issues": validation["issues"],
        "training": validation["training"],
        "location_info": validation["location_info"],
    }
    inspection_records.append(record)
    return jsonify(record)


@app.route("/api/records")
def get_records():
    return jsonify(inspection_records)


@app.route("/api/master_data")
def get_master_data():
    return jsonify(MASTER_DATA)


@app.route("/api/export_excel")
def export_excel():
    wb = openpyxl.Workbook()

    # ---- Sheet 1: Inspection Checklist ----
    ws1 = wb.active
    ws1.title = "Sheet1 - Inspection Checklist"

    # Colors
    header_fill = PatternFill("solid", fgColor="1F4E79")
    pass_fill = PatternFill("solid", fgColor="C6EFCE")
    fail_fill = PatternFill("solid", fgColor="FFC7CE")
    crit_fill = PatternFill("solid", fgColor="FF0000")
    stop_fill = PatternFill("solid", fgColor="FFEB9C")
    alt_fill = PatternFill("solid", fgColor="DCE6F1")

    header_font = Font(name="TH Sarabun New", bold=True, color="FFFFFF", size=12)
    body_font = Font(name="TH Sarabun New", size=11)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    thin = Side(style="thin", color="AAAAAA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Title
    ws1.merge_cells("A1:M1")
    title_cell = ws1["A1"]
    title_cell.value = "SAFETY-EE SYSTEM — Inspection Checklist | Plan B Media Standard"
    title_cell.font = Font(name="TH Sarabun New", bold=True, size=14, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor="1F4E79")
    title_cell.alignment = center
    ws1.row_dimensions[1].height = 30

    headers = [
        "วันที่ตรวจ", "รหัสป้าย\n(Code ID)", "Location",
        "ช่างผู้ตรวจ", "PPE\nครบ?", "บันทึกมือวัด",
        "Ground\n(Ω)", "Continuity\n(Ω)", "Voltage\n(V)",
        "Location\nถูกต้อง?", "อุปกรณ์\nผ่าน?", "ปัญหาพบ", "ผลการตรวจ"
    ]
    ws1.append(headers)
    for col, cell in enumerate(ws1[2], 1):
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
    ws1.row_dimensions[2].height = 35

    col_widths = [16, 12, 20, 16, 10, 14, 10, 12, 10, 14, 12, 35, 14]
    for i, w in enumerate(col_widths, 1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    for idx, rec in enumerate(inspection_records):
        ppe_ok = "✓" if len(rec.get("ppe", [])) == 4 else "✗"
        loc_ok = "✓" if rec.get("code_id") in MASTER_DATA else "✗"
        eq_ok = "✓" if len(rec.get("equipment", [])) == 4 else "✗"
        issues_str = "; ".join(rec.get("issues", [])) if rec.get("issues") else "-"
        row = [
            rec["date"], rec["code_id"], rec["location"],
            rec["technician"], ppe_ok, "มีรูป",
            rec["ground"], rec["continuity"], rec["voltage"],
            loc_ok, eq_ok, issues_str, rec["result"]
        ]
        ws1.append(row)
        row_num = idx + 3
        result = rec["result"]
        row_fill = pass_fill if result == "PASS" else (
            crit_fill if result == "CRITICAL" else (
                stop_fill if result == "STOP" else fail_fill))
        result_font_color = "FFFFFF" if result == "CRITICAL" else "000000"

        for col_idx, cell in enumerate(ws1[row_num], 1):
            cell.font = body_font
            cell.border = border
            cell.alignment = center if col_idx != 12 else left
            if col_idx == 13:
                cell.fill = row_fill
                cell.font = Font(name="TH Sarabun New", bold=True,
                                 size=11, color=result_font_color)
            elif idx % 2 == 1:
                cell.fill = alt_fill
        ws1.row_dimensions[row_num].height = 22

    # Summary row
    total = len(inspection_records)
    passed = sum(1 for r in inspection_records if r["result"] == "PASS")
    failed = sum(1 for r in inspection_records if r["result"] == "FAIL")
    critical = sum(1 for r in inspection_records if r["result"] == "CRITICAL")
    stopped = sum(1 for r in inspection_records if r["result"] == "STOP")
    sum_row = total + 3
    ws1.merge_cells(f"A{sum_row}:K{sum_row}")
    ws1[f"A{sum_row}"].value = f"สรุป: ทั้งหมด {total} รายการ | PASS: {passed} | FAIL: {failed} | CRITICAL: {critical} | STOP: {stopped}"
    ws1[f"A{sum_row}"].font = Font(name="TH Sarabun New", bold=True, size=11)
    ws1[f"A{sum_row}"].fill = PatternFill("solid", fgColor="BDD7EE")
    ws1[f"A{sum_row}"].alignment = center
    ws1.row_dimensions[sum_row].height = 22

    # ---- Sheet 2: Training Matrix ----
    ws2 = wb.create_sheet("Sheet2 - Training Matrix")

    ws2.merge_cells("A1:F1")
    ws2["A1"].value = "Training Matrix — Auto Triggered by SAFETY-EE System"
    ws2["A1"].font = Font(name="TH Sarabun New", bold=True, size=14, color="FFFFFF")
    ws2["A1"].fill = PatternFill("solid", fgColor="7030A0")
    ws2["A1"].alignment = center
    ws2.row_dimensions[1].height = 30

    t_headers = ["ปัญหาที่พบ", "ระดับความเสี่ยง", "หลักสูตรที่ต้องเรียน", "ประเภท", "ระยะเวลา", "Mapping อัตโนมัติ"]
    ws2.append(t_headers)
    purple_fill = PatternFill("solid", fgColor="7030A0")
    for cell in ws2[2]:
        cell.font = header_font
        cell.fill = purple_fill
        cell.alignment = center
        cell.border = border
    ws2.row_dimensions[2].height = 30

    t_col_widths = [28, 18, 28, 14, 14, 22]
    for i, w in enumerate(t_col_widths, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    level_colors = {"สูง": "FFC7CE", "ปานกลาง": "FFEB9C", "วิกฤต": "FF0000", "ต่ำ": "C6EFCE"}
    type_icons = {"New Skill": "★ New Skill", "Reskill": "↺ Reskill", "Upskill": "↑ Upskill"}

    # Collect all training triggers from records
    training_summary = {}
    for rec in inspection_records:
        for tr in rec.get("training", []):
            key = tr["issue"]
            if key not in training_summary:
                training_summary[key] = {**tr, "count": 0, "technicians": []}
            training_summary[key]["count"] += 1
            if rec["technician"] not in training_summary[key]["technicians"]:
                training_summary[key]["technicians"].append(rec["technician"])

    if not training_summary:
        for issue, info in TRAINING_MATRIX.items():
            row = [issue, info["level"], info["course"],
                   type_icons.get(info["type"], info["type"]),
                   info["duration"], "รอ Trigger"]
            ws2.append(row)
            r = ws2.max_row
            lv_fill = PatternFill("solid", fgColor=level_colors.get(info["level"], "FFFFFF"))
            for ci, cell in enumerate(ws2[r], 1):
                cell.font = body_font
                cell.border = border
                cell.alignment = center
                if ci == 2:
                    cell.fill = lv_fill
                elif r % 2 == 0:
                    cell.fill = alt_fill
            ws2.row_dimensions[r].height = 22
    else:
        for issue, info in training_summary.items():
            techs = ", ".join(info["technicians"])
            row = [issue, info["level"], info["course"],
                   type_icons.get(info["type"], info["type"]),
                   info["duration"], f"Triggered {info['count']}x → {techs}"]
            ws2.append(row)
            r = ws2.max_row
            lv_fill = PatternFill("solid", fgColor=level_colors.get(info["level"], "FFFFFF"))
            lv_font_c = "FFFFFF" if info["level"] == "วิกฤต" else "000000"
            for ci, cell in enumerate(ws2[r], 1):
                cell.font = body_font
                cell.border = border
                cell.alignment = center if ci != 6 else left
                if ci == 2:
                    cell.fill = lv_fill
                    cell.font = Font(name="TH Sarabun New", bold=True,
                                     size=11, color=lv_font_c)
                elif r % 2 == 0:
                    cell.fill = alt_fill
            ws2.row_dimensions[r].height = 22

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"SAFETY_EE_Report_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(output, as_attachment=True,
                     download_name=filename,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    app.run(debug=True, port=5050)
